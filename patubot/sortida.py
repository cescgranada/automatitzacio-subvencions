"""Sortides: fitxa Word, pujada a Drive i correu."""
import html
import io
import json
import os
import re
import smtplib
from datetime import date
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText

from . import config

CAMPS_EXTRA = (  # camps que s'afegeixen al final de la fitxa si la plantilla no té el marcador
    ("elegibilitat", "Elegibilitat"), ("requisits", "Requisits clau"), ("adaptacio", "Com adaptar-ho"),
    ("encaix", "Encaix (1-10)"), ("link", "Enllaç"),
)


# ---------------------------------------------------------------- Word
def _valors_fitxa(v: dict) -> dict:
    elig = v.get("elegibilitat", "")
    if v.get("motiu_elegibilitat"):
        elig = f"{elig} — {v['motiu_elegibilitat']}"
    termini = v.get("termini_text") or v.get("termini_data") or "Pendent de publicació"
    if v.get("termini_data") and v["termini_data"] not in termini:
        termini += f" ({v['termini_data']})"
    return {
        "titol": v.get("titol", "-"), "organisme": v.get("organisme", "-"), "import": v.get("import_total", "Desconegut"),
        "termini": termini, "resum": v.get("resum", "-"), "accions": v.get("accions", "-"),
        "adaptacio": v.get("adaptacio") or "-", "elegibilitat": elig or "-", "requisits": v.get("requisits") or "-",
        "encaix": str(v.get("encaix", "-")), "link": v.get("link", "-"),
    }


def _omple_paragraf(p, valors: dict) -> set:
    """Substitueix {{marcador}} encara que Word l'hagi trencat en diversos 'runs'; conserva el format del primer."""
    usats = set()
    for _ in range(20):
        runs = p.runs
        m = re.search(r"\{\{\s*(\w+)\s*\}\}", "".join(r.text for r in runs))
        if not m:
            break
        a, b = m.span()
        pos, primer, ultim = 0, None, None
        for i, r in enumerate(runs):
            fin = pos + len(r.text)
            if primer is None and a < fin:
                primer, inici_primer = i, pos
            if b <= fin:
                ultim = i
                break
            pos = fin
        junts = "".join(r.text for r in runs[primer:ultim + 1])
        runs[primer].text = junts[:a - inici_primer] + valors.get(m.group(1), "-") + junts[b - inici_primer:]
        for r in runs[primer + 1:ultim + 1]:
            r.text = ""
        usats.add(m.group(1))
    return usats


def crea_fitxa_word(v: dict, plantilla="plantilla_subvencio.docx"):
    from docx import Document
    doc = Document(plantilla)
    valors = _valors_fitxa(v)
    usats = set()
    parells = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in parells:
        usats |= _omple_paragraf(p, valors)
    for clau, etiqueta in CAMPS_EXTRA:
        if clau not in usats:
            doc.add_paragraph(f"{etiqueta}: {valors[clau]}")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- Drive
def pujar_a_drive(contingut, nom: str, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"):
    """Llança excepció si falla: el cridador decideix què fer (però ja no es perd en silenci)."""
    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaIoBaseUpload
    creds = service_account.Credentials.from_service_account_info(
        json.loads(os.environ["GDRIVE_CREDENTIALS"]), scopes=["https://www.googleapis.com/auth/drive"])
    servei = build("drive", "v3", credentials=creds, cache_discovery=False)
    servei.files().create(
        body={"name": nom, "parents": [config.GDRIVE_FOLDER_ID]},
        media_body=MediaIoBaseUpload(io.BytesIO(contingut.read()), mimetype=mimetype),
        supportsAllDrives=True,
    ).execute()


def nom_fitxer(v: dict, avui: date) -> str:
    net = re.sub(r"[^\w\- ]+", "", v.get("titol", "subvencio"))[:50].strip()
    return f"E{v.get('encaix', 0):02d}_{avui.isoformat()}_{net}.docx"


# ---------------------------------------------------------------- Correu
def _bloc_text(v: dict) -> str:
    l = [f"[ENCAIX {v['encaix']}/10 · {v['elegibilitat']}] {v['titol']}",
         f"Organisme: {v['organisme']}   Import: {v['import_total']}",
         f"Termini: {v['termini_text'] or 'Pendent'}" + (f" ({v['termini_data']})" if v['termini_data'] else ""),
         f"Resum: {v['resum']}", f"Requisits: {v['requisits']}", f"Primer pas: {v['accions']}"]
    if v.get("adaptacio"):
        l.append(f"Com adaptar-ho: {v['adaptacio']}")
    if v.get("motiu_elegibilitat"):
        l.append(f"Elegibilitat: {v['motiu_elegibilitat']}")
    l.append(f"Enllaç: {v['link']}")
    return "\n".join(l)


def _bloc_html(v: dict) -> str:
    e = html.escape
    color = "#1a7f37" if v["encaix"] >= 8 else "#9a6700"
    termini = e(v["termini_text"] or "Pendent") + (f" <b>({e(v['termini_data'])})</b>" if v["termini_data"] else "")
    extra = f"<p><b>Com adaptar-ho:</b> {e(v['adaptacio'])}</p>" if v.get("adaptacio") else ""
    return (f'<div style="border-left:4px solid {color};padding:6px 12px;margin:14px 0">'
            f'<h3 style="margin:0 0 4px"><span style="color:{color}">Encaix {v["encaix"]}/10</span> · {e(v["titol"])}</h3>'
            f'<p style="margin:2px 0;color:#555">{e(v["organisme"])} · Import: {e(v["import_total"])} · Elegibilitat: {e(v["elegibilitat"])}</p>'
            f'<p><b>Termini:</b> {termini}</p><p>{e(v["resum"])}</p>'
            f'<p><b>Requisits:</b> {e(v["requisits"])}</p><p><b>Primer pas:</b> {e(v["accions"])}</p>{extra}'
            f'<p><a href="{e(v["link"])}">Obrir la convocatòria</a></p></div>')


def construeix_correu(oportunitats, recordatoris, errors, resum: str, avui: date):
    """Retorna (assumpte, text pla, html)."""
    dia = avui.strftime("%d/%m/%Y")
    if oportunitats:
        assumpte = f"Patu-Bot: {len(oportunitats)} oportunitat(s) nova(es) — {dia}"
    elif recordatoris:
        assumpte = f"Patu-Bot: recordatori de termini — {dia}"
    elif errors:
        assumpte = f"Patu-Bot: ATENCIÓ, fonts amb errors — {dia}"
    else:
        assumpte = f"Patu-Bot: resum setmanal, cap novetat — {dia}"

    text, h = [resum, ""], [f"<p>{html.escape(resum)}</p>"]
    if recordatoris:
        text.append("== TERMINIS PROPERS ==")
        h.append("<h2>⏰ Terminis propers</h2><ul>")
        for o in recordatoris:
            text.append(f"- {o['termini_data']}: {o['titol']} — {o['link']}")
            h.append(f'<li><b>{html.escape(o["termini_data"])}</b>: <a href="{html.escape(o["link"] or "")}">{html.escape(o["titol"])}</a></li>')
        h.append("</ul>")
        text.append("")
    if oportunitats:
        text.append("== OPORTUNITATS NOVES ==")
        h.append("<h2>🌟 Oportunitats noves</h2>")
        for v in oportunitats:
            text += [_bloc_text(v), "-" * 60]
            h.append(_bloc_html(v))
    if errors:
        text += ["", "== FONTS AMB ERROR =="] + [f"- {e}" for e in errors]
        h.append("<h3>⚠️ Fonts amb error</h3><ul>" + "".join(f"<li>{html.escape(e)}</li>" for e in errors) + "</ul>")
    return assumpte, "\n".join(text), "<html><body style='font-family:sans-serif'>" + "".join(h) + "</body></html>"


def enviar_correu(assumpte: str, text: str, cos_html: str):
    u, p, r = os.environ["EMAIL_USER"], os.environ["EMAIL_PASS"], os.environ["EMAIL_RECEIVER"]
    msg = MIMEMultipart("alternative")
    msg["Subject"], msg["From"], msg["To"] = assumpte, u, r
    msg.attach(MIMEText(text, "plain", "utf-8"))
    msg.attach(MIMEText(cos_html, "html", "utf-8"))
    with smtplib.SMTP_SSL("smtp.gmail.com", 465, timeout=60) as s:
        s.login(u, p)
        s.sendmail(u, [x.strip() for x in r.split(",")], msg.as_string())
